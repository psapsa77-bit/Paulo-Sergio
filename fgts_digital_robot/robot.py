"""
Classe principal do FGTS Digital Robot.
Orquestra autenticação, navegação, extração e exportação.
"""

from __future__ import annotations

import json
import logging
import time
from pathlib import Path
from typing import List, Optional

from .authenticator import CertificadoError, FGTSAuthenticator
from .extractor import FGTSExtractor
from .models import (
    CertificadoDigital,
    EmpresaFGTS,
    GuiaFGTS,
    RelatorioFGTS,
    StatusGuia,
)
from .navigator import FGTSNavigator, NavegacaoError

logger = logging.getLogger(__name__)


class FGTSRobot:
    """
    Robô de automação para o portal FGTS Digital.

    Autenticação, listagem de empresas, extração de guias e geração de relatórios.

    Exemplo de uso básico::

        robot = FGTSRobot("certificado.pfx", "senha123")
        robot.autenticar()
        relatorio = robot.processar_todas_empresas()
        robot.exportar_relatorio(relatorio, "relatorio.json")
        robot.fechar()

    Ou com context manager::

        with FGTSRobot("certificado.pfx", "senha123") as robot:
            robot.autenticar()
            relatorio = robot.processar_todas_empresas()
    """

    def __init__(
        self,
        certificado_path: str | Path,
        senha_certificado: str,
        headless: bool = False,
        log_level: str = "INFO",
    ):
        """
        Args:
            certificado_path: Caminho para o arquivo .pfx ou .p12
            senha_certificado: Senha do certificado digital
            headless: Se True, o navegador roda sem interface gráfica
            log_level: Nível de log (DEBUG, INFO, WARNING, ERROR)
        """
        logging.basicConfig(
            level=getattr(logging, log_level.upper(), logging.INFO),
            format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
        )

        self.autenticador = FGTSAuthenticator(certificado_path, senha_certificado)
        self.navigator = FGTSNavigator(self.autenticador, headless=headless)
        self.extractor = FGTSExtractor()

        self._autenticado = False
        self._empresas: List[EmpresaFGTS] = []

    # ------------------------------------------------------------------
    # Context manager
    # ------------------------------------------------------------------

    def __enter__(self) -> "FGTSRobot":
        return self

    def __exit__(self, *args) -> None:
        self.fechar()

    # ------------------------------------------------------------------
    # Autenticação
    # ------------------------------------------------------------------

    def validar_certificado(self) -> CertificadoDigital:
        """
        Valida o certificado digital sem iniciar o browser.

        Returns:
            CertificadoDigital com informações do certificado.

        Raises:
            CertificadoError: Se o certificado for inválido ou vencido.
        """
        return self.autenticador.validar()

    def autenticar(self) -> None:
        """
        Inicia o browser, valida o certificado e faz login no portal.

        Raises:
            CertificadoError: Se o certificado for inválido.
            NavegacaoError: Se o login falhar.
        """
        logger.info("Iniciando autenticação…")

        # Valida certificado antes de abrir browser
        cert = self.autenticador.validar()
        logger.info(
            "Certificado: %s | Válido até: %s",
            cert.nome_titular or "desconhecido",
            cert.valido_ate,
        )

        if cert.dias_para_vencer is not None and cert.dias_para_vencer < 30:
            logger.warning(
                "ATENÇÃO: Seu certificado vence em %d dias! Renove antes de usar.",
                cert.dias_para_vencer,
            )

        self.navigator.iniciar()
        self.navigator.fazer_login()
        self._autenticado = True
        logger.info("Autenticação concluída.")

    # ------------------------------------------------------------------
    # Empresas
    # ------------------------------------------------------------------

    def listar_empresas(self) -> List[EmpresaFGTS]:
        """
        Retorna lista de empresas disponíveis para consulta.

        Returns:
            Lista de EmpresaFGTS com procuração ativa.
        """
        self._verificar_autenticado()
        self._empresas = self.navigator.listar_empresas()
        return self._empresas

    # ------------------------------------------------------------------
    # Guias
    # ------------------------------------------------------------------

    def extrair_guias_empresa(
        self,
        cnpj: str,
        filtrar_status: Optional[List[StatusGuia]] = None,
    ) -> List[GuiaFGTS]:
        """
        Extrai guias de uma empresa específica.

        Args:
            cnpj: CNPJ da empresa (formatado ou só dígitos)
            filtrar_status: Se fornecido, retorna apenas guias com esses status

        Returns:
            Lista de GuiaFGTS
        """
        self._verificar_autenticado()

        logger.info("Extraindo guias da empresa CNPJ: %s", cnpj)

        sucesso = self.navigator.selecionar_empresa(cnpj)
        if not sucesso:
            logger.warning("Não foi possível selecionar empresa %s. Tentando navegar direto.", cnpj)

        self.navigator.navegar_para_guias()
        html = self.navigator.obter_html_guias()
        guias = self.extractor.extrair_guias(html, cnpj)

        if filtrar_status:
            guias = [g for g in guias if g.status in filtrar_status]

        logger.info("Extraídas %d guias para CNPJ %s.", len(guias), cnpj)
        return guias

    def processar_todas_empresas(
        self,
        filtrar_status: Optional[List[StatusGuia]] = None,
    ) -> RelatorioFGTS:
        """
        Processa todas as empresas disponíveis e gera relatório consolidado.

        Args:
            filtrar_status: Lista de status para filtrar guias. None = todos.

        Returns:
            RelatorioFGTS com todas as guias e estatísticas.
        """
        self._verificar_autenticado()

        inicio = time.time()
        relatorio = RelatorioFGTS()

        # Listar empresas se ainda não foi feito
        if not self._empresas:
            self._empresas = self.listar_empresas()

        relatorio.empresas = list(self._empresas)

        if not self._empresas:
            logger.warning(
                "Nenhuma empresa encontrada. Verifique se há procurações ativas no portal."
            )
            return relatorio

        logger.info("Processando %d empresa(s)…", len(self._empresas))

        for i, empresa in enumerate(self._empresas, 1):
            logger.info("[%d/%d] Empresa: %s (%s)", i, len(self._empresas), empresa.razao_social, empresa.cnpj)
            try:
                guias = self.extrair_guias_empresa(empresa.cnpj, filtrar_status)
                relatorio.guias.extend(guias)
            except Exception as exc:
                logger.error("Erro ao processar empresa %s: %s", empresa.cnpj, exc)

        relatorio.tempo_execucao_segundos = time.time() - inicio

        logger.info(
            "Processamento concluído em %.1fs | Total: %d guias | Pendentes: %d | Valor pendente: R$ %s",
            relatorio.tempo_execucao_segundos,
            relatorio.total_guias,
            relatorio.total_pendentes,
            relatorio.valor_total_pendente,
        )

        return relatorio

    def processar_empresas(
        self,
        cnpjs: List[str],
        filtrar_status: Optional[List[StatusGuia]] = None,
    ) -> RelatorioFGTS:
        """
        Processa uma lista específica de CNPJs.

        Args:
            cnpjs: Lista de CNPJs a processar
            filtrar_status: Filtro de status

        Returns:
            RelatorioFGTS
        """
        self._verificar_autenticado()

        inicio = time.time()
        relatorio = RelatorioFGTS()

        for cnpj in cnpjs:
            empresa = self._encontrar_empresa(cnpj)
            if empresa:
                relatorio.empresas.append(empresa)
            try:
                guias = self.extrair_guias_empresa(cnpj, filtrar_status)
                relatorio.guias.extend(guias)
            except Exception as exc:
                logger.error("Erro ao processar CNPJ %s: %s", cnpj, exc)

        relatorio.tempo_execucao_segundos = time.time() - inicio
        return relatorio

    # ------------------------------------------------------------------
    # Exportação
    # ------------------------------------------------------------------

    def exportar_relatorio(
        self,
        relatorio: RelatorioFGTS,
        caminho_saida: str | Path,
        formato: str = "json",
    ) -> Path:
        """
        Exporta o relatório para arquivo.

        Args:
            relatorio: RelatorioFGTS gerado
            caminho_saida: Caminho do arquivo de saída
            formato: 'json' ou 'excel'

        Returns:
            Path do arquivo gerado.
        """
        caminho = Path(caminho_saida)

        if formato.lower() == "excel":
            return self._exportar_excel(relatorio, caminho)
        else:
            return self._exportar_json(relatorio, caminho)

    def _exportar_json(self, relatorio: RelatorioFGTS, caminho: Path) -> Path:
        if not caminho.suffix:
            caminho = caminho.with_suffix(".json")
        caminho.parent.mkdir(parents=True, exist_ok=True)

        with open(caminho, "w", encoding="utf-8") as f:
            json.dump(relatorio.to_dict(), f, ensure_ascii=False, indent=2)

        logger.info("Relatório JSON salvo: %s", caminho)
        return caminho

    def _exportar_excel(self, relatorio: RelatorioFGTS, caminho: Path) -> Path:
        try:
            import openpyxl
        except ImportError as exc:
            raise ImportError(
                "openpyxl não instalado. Execute: pip install openpyxl"
            ) from exc

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        if not caminho.suffix or caminho.suffix.lower() != ".xlsx":
            caminho = caminho.with_suffix(".xlsx")
        caminho.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()

        # Aba Resumo
        ws_resumo = wb.active
        ws_resumo.title = "Resumo"
        self._preencher_aba_resumo(ws_resumo, relatorio)

        # Aba Guias
        ws_guias = wb.create_sheet("Guias")
        self._preencher_aba_guias(ws_guias, relatorio)

        # Aba Empresas
        ws_emp = wb.create_sheet("Empresas")
        self._preencher_aba_empresas(ws_emp, relatorio)

        wb.save(caminho)
        logger.info("Relatório Excel salvo: %s", caminho)
        return caminho

    def _preencher_aba_resumo(self, ws, relatorio: RelatorioFGTS) -> None:
        from openpyxl.styles import Font
        ws["A1"] = "Relatório FGTS Digital"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A3"] = "Gerado em"
        ws["B3"] = relatorio.gerado_em.strftime("%d/%m/%Y %H:%M")
        ws["A4"] = "Tempo de execução (s)"
        ws["B4"] = round(relatorio.tempo_execucao_segundos, 1)
        ws["A6"] = "Total de guias"
        ws["B6"] = relatorio.total_guias
        ws["A7"] = "Pendentes"
        ws["B7"] = relatorio.total_pendentes
        ws["A8"] = "Pagas"
        ws["B8"] = relatorio.total_pagas
        ws["A9"] = "Vencidas"
        ws["B9"] = relatorio.total_vencidas
        ws["A10"] = "Parceladas"
        ws["B10"] = relatorio.total_parceladas
        ws["A12"] = "Valor pendente (R$)"
        ws["B12"] = float(relatorio.valor_total_pendente)
        ws["A13"] = "Valor pago (R$)"
        ws["B13"] = float(relatorio.valor_total_pago)

    def _preencher_aba_guias(self, ws, relatorio: RelatorioFGTS) -> None:
        from openpyxl.styles import Font, PatternFill
        cabecalho = [
            "Número Guia", "CNPJ Empresa", "Tipo", "Status", "Competência",
            "Vencimento", "Pagamento", "Valor Principal", "Valor Multa",
            "Valor Juros", "Valor Total", "Dias Atraso",
        ]
        ws.append(cabecalho)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        for guia in relatorio.guias:
            ws.append([
                guia.numero_guia,
                guia.cnpj_empresa,
                guia.tipo.value,
                guia.status.value,
                guia.competencia,
                guia.data_vencimento.strftime("%d/%m/%Y") if guia.data_vencimento else "",
                guia.data_pagamento.strftime("%d/%m/%Y") if guia.data_pagamento else "",
                float(guia.valor_principal),
                float(guia.valor_multa),
                float(guia.valor_juros),
                float(guia.valor_total),
                guia.dias_atraso or 0,
            ])

    def _preencher_aba_empresas(self, ws, relatorio: RelatorioFGTS) -> None:
        from openpyxl.styles import Font
        cabecalho = ["CNPJ", "Razão Social", "Nome Fantasia", "Total Guias", "Pendentes", "Valor Pendente"]
        ws.append(cabecalho)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        for empresa in relatorio.empresas:
            guias_emp = relatorio.guias_por_empresa(empresa.cnpj)
            pendentes = [g for g in guias_emp if g.status in (StatusGuia.PENDENTE, StatusGuia.VENCIDA)]
            valor_pend = sum((g.valor_total for g in pendentes), __import__("decimal").Decimal("0"))
            ws.append([
                empresa.cnpj_formatado,
                empresa.razao_social,
                empresa.nome_fantasia,
                len(guias_emp),
                len(pendentes),
                float(valor_pend),
            ])

    # ------------------------------------------------------------------
    # Utilitários
    # ------------------------------------------------------------------

    def fazer_screenshot(self, caminho: str) -> None:
        """Salva screenshot da tela atual para debug."""
        self.navigator.fazer_screenshot(caminho)

    def fechar(self) -> None:
        """Fecha o browser e libera recursos."""
        self.navigator.fechar()
        self._autenticado = False
        logger.info("Robô encerrado.")

    def _verificar_autenticado(self) -> None:
        if not self._autenticado:
            raise RuntimeError("Robô não autenticado. Chame autenticar() primeiro.")

    def _encontrar_empresa(self, cnpj: str) -> Optional[EmpresaFGTS]:
        cnpj_limpo = cnpj.replace(".", "").replace("/", "").replace("-", "")
        for e in self._empresas:
            if e.cnpj.replace(".", "").replace("/", "").replace("-", "") == cnpj_limpo:
                return e
        return None
