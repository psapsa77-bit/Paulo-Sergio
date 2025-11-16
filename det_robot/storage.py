"""
Módulo de armazenamento para dados extraídos do DET
"""

import json
import csv
from datetime import datetime
from pathlib import Path
from typing import Optional, List

from .models import ResultadoExtracao, MensagemDET


class DETStorage:
    """Gerencia o armazenamento dos dados extraídos"""

    def __init__(self, pasta_base: str = "./dados_det", log_callback=None):
        """
        Inicializa o armazenamento

        Args:
            pasta_base: Pasta base para salvar dados
            log_callback: Função para logging
        """
        self.pasta_base = Path(pasta_base)
        self.pasta_base.mkdir(parents=True, exist_ok=True)
        self.log = log_callback or print

    def _gerar_nome_arquivo(self, prefixo: str, extensao: str) -> Path:
        """Gera nome de arquivo com timestamp"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return self.pasta_base / f"{prefixo}_{timestamp}.{extensao}"

    def salvar_json(
        self,
        resultado: ResultadoExtracao,
        nome_arquivo: Optional[str] = None
    ) -> str:
        """
        Salva resultado em formato JSON

        Args:
            resultado: ResultadoExtracao a salvar
            nome_arquivo: Nome do arquivo (opcional)

        Returns:
            Caminho do arquivo salvo
        """
        try:
            if nome_arquivo:
                caminho = self.pasta_base / nome_arquivo
            else:
                caminho = self._gerar_nome_arquivo("extracao_det", "json")

            # Converter para dicionário serializável
            dados = resultado.model_dump(mode="json")

            with open(caminho, "w", encoding="utf-8") as f:
                json.dump(dados, f, ensure_ascii=False, indent=2, default=str)

            self.log(f"Dados salvos em JSON: {caminho}")
            return str(caminho)

        except Exception as e:
            self.log(f"Erro ao salvar JSON: {e}")
            raise

    def salvar_csv(
        self,
        resultado: ResultadoExtracao,
        nome_arquivo: Optional[str] = None
    ) -> str:
        """
        Salva mensagens em formato CSV

        Args:
            resultado: ResultadoExtracao a salvar
            nome_arquivo: Nome do arquivo (opcional)

        Returns:
            Caminho do arquivo salvo
        """
        try:
            if nome_arquivo:
                caminho = self.pasta_base / nome_arquivo
            else:
                caminho = self._gerar_nome_arquivo("mensagens_det", "csv")

            # Campos do CSV
            campos = [
                "id_mensagem",
                "numero",
                "assunto",
                "tipo",
                "data_envio",
                "data_ciencia",
                "prazo_resposta",
                "dias_restantes",
                "remetente",
                "status",
                "resumo",
                "url_mensagem",
                "anexos",
            ]

            with open(caminho, "w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=campos)
                writer.writeheader()

                for mensagem in resultado.mensagens:
                    row = mensagem.model_dump()
                    # Converter listas e datas para strings
                    row["data_envio"] = row["data_envio"].isoformat() if row["data_envio"] else ""
                    row["data_ciencia"] = row["data_ciencia"].isoformat() if row["data_ciencia"] else ""
                    row["prazo_resposta"] = row["prazo_resposta"].isoformat() if row["prazo_resposta"] else ""
                    row["anexos"] = ";".join(row["anexos"]) if row["anexos"] else ""
                    row["tipo"] = row["tipo"].value if hasattr(row["tipo"], "value") else str(row["tipo"])
                    row["status"] = row["status"].value if hasattr(row["status"], "value") else str(row["status"])
                    writer.writerow(row)

            self.log(f"Dados salvos em CSV: {caminho}")
            return str(caminho)

        except Exception as e:
            self.log(f"Erro ao salvar CSV: {e}")
            raise

    def salvar_excel(
        self,
        resultado: ResultadoExtracao,
        nome_arquivo: Optional[str] = None
    ) -> str:
        """
        Salva mensagens em formato Excel

        Args:
            resultado: ResultadoExtracao a salvar
            nome_arquivo: Nome do arquivo (opcional)

        Returns:
            Caminho do arquivo salvo
        """
        try:
            # Tentar importar openpyxl
            try:
                import openpyxl
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill
            except ImportError:
                self.log("openpyxl não instalado. Salvando como CSV...")
                return self.salvar_csv(resultado, nome_arquivo.replace(".xlsx", ".csv") if nome_arquivo else None)

            if nome_arquivo:
                caminho = self.pasta_base / nome_arquivo
            else:
                caminho = self._gerar_nome_arquivo("mensagens_det", "xlsx")

            wb = Workbook()
            ws = wb.active
            ws.title = "Mensagens DET"

            # Cabeçalho
            headers = [
                "ID",
                "Número",
                "Assunto",
                "Tipo",
                "Data Envio",
                "Data Ciência",
                "Prazo Resposta",
                "Dias Restantes",
                "Remetente",
                "Status",
                "Resumo",
                "URL",
            ]

            # Estilo do cabeçalho
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)

            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font

            # Dados
            for row_idx, mensagem in enumerate(resultado.mensagens, 2):
                ws.cell(row=row_idx, column=1, value=mensagem.id_mensagem)
                ws.cell(row=row_idx, column=2, value=mensagem.numero or "")
                ws.cell(row=row_idx, column=3, value=mensagem.assunto)
                ws.cell(row=row_idx, column=4, value=mensagem.tipo.value)
                ws.cell(row=row_idx, column=5, value=mensagem.data_envio.strftime("%d/%m/%Y %H:%M"))
                ws.cell(row=row_idx, column=6, value=mensagem.data_ciencia.strftime("%d/%m/%Y %H:%M") if mensagem.data_ciencia else "")
                ws.cell(row=row_idx, column=7, value=mensagem.prazo_resposta.strftime("%d/%m/%Y") if mensagem.prazo_resposta else "")
                ws.cell(row=row_idx, column=8, value=mensagem.dias_restantes or "")
                ws.cell(row=row_idx, column=9, value=mensagem.remetente)
                ws.cell(row=row_idx, column=10, value=mensagem.status.value)
                ws.cell(row=row_idx, column=11, value=mensagem.resumo or "")
                ws.cell(row=row_idx, column=12, value=mensagem.url_mensagem or "")

                # Destacar mensagens urgentes (menos de 5 dias)
                if mensagem.dias_restantes is not None and mensagem.dias_restantes < 5:
                    urgente_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                    for col in range(1, 13):
                        ws.cell(row=row_idx, column=col).fill = urgente_fill

            # Ajustar largura das colunas
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column].width = adjusted_width

            wb.save(caminho)
            self.log(f"Dados salvos em Excel: {caminho}")
            return str(caminho)

        except Exception as e:
            self.log(f"Erro ao salvar Excel: {e}")
            raise

    def carregar_json(self, caminho: str) -> ResultadoExtracao:
        """
        Carrega resultado de arquivo JSON

        Args:
            caminho: Caminho do arquivo

        Returns:
            ResultadoExtracao
        """
        try:
            with open(caminho, "r", encoding="utf-8") as f:
                dados = json.load(f)

            resultado = ResultadoExtracao(**dados)
            self.log(f"Dados carregados de: {caminho}")
            return resultado

        except Exception as e:
            self.log(f"Erro ao carregar JSON: {e}")
            raise

    def listar_extracoes(self) -> List[dict]:
        """
        Lista todas as extrações salvas

        Returns:
            Lista de dicionários com informações dos arquivos
        """
        arquivos = []

        for arquivo in self.pasta_base.glob("extracao_det_*.json"):
            try:
                stat = arquivo.stat()
                arquivos.append({
                    "nome": arquivo.name,
                    "caminho": str(arquivo),
                    "tamanho": stat.st_size,
                    "data_modificacao": datetime.fromtimestamp(stat.st_mtime),
                })
            except Exception as e:
                self.log(f"Erro ao ler {arquivo}: {e}")

        # Ordenar por data (mais recente primeiro)
        arquivos.sort(key=lambda x: x["data_modificacao"], reverse=True)
        return arquivos

    def gerar_relatorio_texto(self, resultado: ResultadoExtracao) -> str:
        """
        Gera relatório em texto formatado

        Args:
            resultado: ResultadoExtracao

        Returns:
            String com relatório formatado
        """
        linhas = []
        linhas.append("=" * 60)
        linhas.append("RELATÓRIO DE EXTRAÇÃO - DET")
        linhas.append("=" * 60)
        linhas.append(f"Data: {resultado.data_extracao.strftime('%d/%m/%Y %H:%M')}")
        linhas.append(f"Total de mensagens: {resultado.total_extraido}")
        linhas.append(f"Apenas não lidas: {'Sim' if resultado.apenas_nao_lidas else 'Não'}")

        if resultado.tempo_execucao:
            linhas.append(f"Tempo de execução: {resultado.tempo_execucao:.2f} segundos")

        if resultado.empregador:
            linhas.append("\n--- DADOS DO EMPREGADOR ---")
            linhas.append(f"CNPJ/CPF: {resultado.empregador.cnpj_cpf}")
            if resultado.empregador.razao_social:
                linhas.append(f"Razão Social: {resultado.empregador.razao_social}")
            linhas.append(f"Mensagens não lidas: {resultado.empregador.mensagens_nao_lidas}")

        linhas.append("\n--- MENSAGENS ---")
        for i, msg in enumerate(resultado.mensagens, 1):
            linhas.append(f"\n{i}. {msg.assunto}")
            linhas.append(f"   Tipo: {msg.tipo.value}")
            linhas.append(f"   Data: {msg.data_envio.strftime('%d/%m/%Y')}")
            linhas.append(f"   Status: {msg.status.value}")
            if msg.prazo_resposta:
                linhas.append(f"   Prazo: {msg.prazo_resposta.strftime('%d/%m/%Y')}")
                if msg.dias_restantes is not None:
                    linhas.append(f"   Dias restantes: {msg.dias_restantes}")

        if resultado.erros:
            linhas.append("\n--- ERROS ---")
            for erro in resultado.erros:
                linhas.append(f"  - {erro}")

        linhas.append("\n" + "=" * 60)

        return "\n".join(linhas)

    def salvar_relatorio_texto(
        self,
        resultado: ResultadoExtracao,
        nome_arquivo: Optional[str] = None
    ) -> str:
        """
        Salva relatório em formato texto

        Args:
            resultado: ResultadoExtracao
            nome_arquivo: Nome do arquivo (opcional)

        Returns:
            Caminho do arquivo salvo
        """
        try:
            if nome_arquivo:
                caminho = self.pasta_base / nome_arquivo
            else:
                caminho = self._gerar_nome_arquivo("relatorio_det", "txt")

            relatorio = self.gerar_relatorio_texto(resultado)

            with open(caminho, "w", encoding="utf-8") as f:
                f.write(relatorio)

            self.log(f"Relatório salvo em: {caminho}")
            return str(caminho)

        except Exception as e:
            self.log(f"Erro ao salvar relatório: {e}")
            raise
